import difflib
from docx import Document
from redlines import Redlines
from html.parser import HTMLParser

old_path = 'test-EA-letter.docx'
new_path = 'test-EA-letterRev.docx'
output_path = 'diff.docx'

class DiffHTMLParser(HTMLParser):
    def __init__(self, paragraph):
        super().__init__()
        self.paragraph = paragraph
        self._strike = False
        self._underline = False

    def handle_starttag(self, tag, attrs):
        if tag == 'del':
            self._strike = True
        elif tag == 'ins':
            self._underline = True

    def handle_endtag(self, tag):
        if tag == 'del':
            self._strike = False
        elif tag == 'ins':
            self._underline = False

    def handle_data(self, data):
        run = self.paragraph.add_run(data)
        run.font.strike = self._strike
        run.font.underline = self._underline

def extract_paragraphs(word_file_path):
    doc = Document(word_file_path)
    return [p.text for p in doc.paragraphs]


def diff_paragraphs(old_paragraphs, new_paragraphs):
    matcher = difflib.SequenceMatcher(None, old_paragraphs, new_paragraphs)
    results = []
    for tag, i1, i2, j1, j2 in matcher.get_opcodes():
        old_block = old_paragraphs[i1:i2]
        new_block = new_paragraphs[j1:j2]
        results.append((tag, old_block, new_block))
    return results

def write_diff_docx(diff_results, output_path):
    doc = Document()
    for tag, old_block, new_block in diff_results:
        if tag == 'equal':
            for text in new_block:
                doc.add_paragraph(text)
        elif tag == 'replace':
            for old, new in zip(old_block, new_block):
                diff = Redlines(old, new)
                para = doc.add_paragraph()
                DiffHTMLParser(para).feed(diff.output_markdown)
            for text in old_block[len(new_block):]:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.strike = True
            for text in new_block[len(old_block):]:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.underline = True
        elif tag == 'insert':
            for text in new_block:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.underline = True
        elif tag == 'delete':
            for text in old_block:
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.strike = True
    doc.save(output_path)


def run(old_path, new_path, output_path):
    old_graphs = extract_paragraphs(old_path)
    new_graphs = extract_paragraphs(new_path)
    write_diff_docx(diff_paragraphs(old_graphs, new_graphs), output_path)


if __name__ == "__main__":
    run(old_path, new_path, output_path)
