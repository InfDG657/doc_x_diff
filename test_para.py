from docx import Document
from redlines import Redlines
from html.parser import HTMLParser

old = "The quick brown fox jumps over the lazy dog."
new = "The quick brown fox leaps over the energetic dog."

output_path = "test_para_diff.docx"


class DiffHTMLParser(HTMLParser):
    def __init__(self, paragraph):
        super().__init__()
        self.paragraph = paragraph
        self._strike = False
        self._underline = False

    def handle_starttag(self, tag, attrs):
        if tag == 'del':
            self._strike = True
        elif tag == 'ins':
            self._underline = True

    def handle_endtag(self, tag):
        if tag == 'del':
            self._strike = False
        elif tag == 'ins':
            self._underline = False

    def handle_data(self, data):
        run = self.paragraph.add_run(data)
        run.font.strike = self._strike
        run.font.underline = self._underline


diff = Redlines(old, new)
doc = Document()
para = doc.add_paragraph()
DiffHTMLParser(para).feed(diff.output_html)
doc.save(output_path)
