from docx import Document
from docx.shared import RGBColor
from html.parser import HTMLParser
from redlines import Redlines

PARA_SEP = "\n\n"

class DiffHTMLParser(HTMLParser):
    def __init__(self, paragraph):
        super().__init__()
        self.paragraph = paragraph
        self._strike = False
        self._underline = False

    def handle_starttag(self, tag, attrs):
        if tag == 'del':
            self._strike = True
        elif tag == 'ins':
            self._underline = True

    def handle_endtag(self, tag):
        if tag == 'del':
            self._strike = False
        elif tag == 'ins':
            self._underline = False

    def handle_data(self, data):
        run = self.paragraph.add_run(data)
        run.font.strike = self._strike
        run.font.underline = self._underline
        if self._strike:
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
        elif self._underline:
            run.font.color.rgb = RGBColor(0x00, 0x80, 0x00)


def extract_text(word_file_path):
    doc = Document(word_file_path)
    return PARA_SEP.join([p.text for p in doc.paragraphs])


def write_diff_docx(diff_output, output_path):
    doc = Document()
    for para_html in diff_output.output_markdown.split(PARA_SEP):
        para = doc.add_paragraph()
        DiffHTMLParser(para).feed(para_html)
    doc.save(output_path)

def run(old_path, new_path, output_path):
    diff =  Redlines(extract_text(old_path), 
                     extract_text(new_path), 
                     markdown_style = None)
    write_diff_docx(diff, output_path)

if __name__ == "__main__":
    run('test-EA-letter.docx', 'test-EA-letterRev.docx', 'diff.docx')
